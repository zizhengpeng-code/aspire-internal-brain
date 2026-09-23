"""Human review of eval cases via an Excel sheet (needs openpyxl; the rest of the repo does not).

    python review.py export   # data/eval.json -> review/eval_review.xlsx
    python review.py apply    # review/eval_review.xlsx -> data/eval.json

In the sheet, fill the 确认 column per row:
    ✓  confirm as-is                  -> source = "user"
    改 confirm with the edits you typed in 修改后问题 / 修改后 gold -> source = "user"
    删 remove the case
    (blank) leave as draft
"""
import json, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

CORPUS, EVAL, SHEET = "data/corpus.json", "data/eval.json", "review/eval_review.xlsx"
COLS = [("ID", 7), ("角色", 16), ("类型", 11), ("期望", 8), ("问题", 45), ("gold", 11),
        ("gold 标题", 28), ("gold 原文", 70), ("确认", 7), ("修改后问题", 30), ("修改后 gold", 12), ("备注", 25)]
TAG_ZH = {"plain": "普通", "colloquial": "口语", "zh_mix": "中英混合", "cross_doc": "跨文档", "injection": "注入"}
EXPECT_ZH = {"answer": "应回答", "deny": "应拒答"}

def load():
    docs = {d["id"]: d for d in json.load(open(CORPUS, encoding="utf-8"))["documents"]}
    return docs, json.load(open(EVAL, encoding="utf-8"))

def export():
    docs, cases = load()
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    ws.append([c for c, _ in COLS])
    for i, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in cases:
        gold = [c["gold"]] if isinstance(c["gold"], str) else c["gold"]
        ws.append([c["id"], c["role"], TAG_ZH.get(c.get("tag"), c.get("tag")), EXPECT_ZH[c["expect"]],
                   c["question"], ", ".join(gold),
                   "\n".join(docs[g]["title"] for g in gold),
                   "\n\n".join(f"[{g}] {docs[g]['text']}" for g in gold),
                   "✓" if c.get("source") == "user" else "", "", "", ""])
    head = PatternFill("solid", fgColor="DDE4EE")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), head
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[3].value == "应拒答":
            row[3].fill = PatternFill("solid", fgColor="F8E0E0")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1='"✓,改,删"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{ws.max_row}")

    ds = wb.create_sheet("文档")
    ds.append(["ID", "标题", "部门", "密级", "可见角色", "原文"])
    for col, w in zip("ABCDEF", (7, 30, 12, 12, 30, 90)):
        ds.column_dimensions[col].width = w
    for d in docs.values():
        ds.append([d["id"], d["title"], d["department"], d["sensitivity"], ", ".join(d["allowed_roles"]), d["text"]])
    for cell in ds[1]:
        cell.font, cell.fill = Font(bold=True), head
    for row in ds.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ds.freeze_panes = "C2"

    import os
    os.makedirs("review", exist_ok=True)
    wb.save(SHEET)
    print(f"wrote {SHEET}: {len(cases)} cases, {len(docs)} docs")

def apply():
    docs, cases = load()
    ws = load_workbook(SHEET)["测试用例"]
    idx = {name: i for i, (name, _) in enumerate(COLS)}
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    out, n = [], {"✓": 0, "改": 0, "删": 0, "": 0}
    for c in cases:
        r = rows.get(c["id"])
        mark = (r[idx["确认"]] or "").strip() if r else ""
        if mark not in n:
            raise SystemExit(f"{c['id']}: unknown mark {mark!r} (use ✓ / 改 / 删 or leave blank)")
        n[mark] += 1
        if mark == "删":
            continue
        if mark == "改":
            q, g = r[idx["修改后问题"]], r[idx["修改后 gold"]]
            if q:
                c["question"] = str(q).strip()
            if g:
                ids = [x.strip().upper() for x in str(g).replace("，", ",").split(",") if x.strip()]
                unknown = [x for x in ids if x not in docs]
                if unknown:
                    raise SystemExit(f"{c['id']}: unknown doc id {unknown}")
                c["gold"] = ids[0] if len(ids) == 1 else ids
        if mark in ("✓", "改"):
            c["source"] = "user"
        if r and r[idx["备注"]]:
            c["note"] = str(r[idx["备注"]]).strip()
        out.append(c)
    with open(EVAL, "w", encoding="utf-8") as f:
        f.write("[\n" + ",\n".join(" " + json.dumps(c, ensure_ascii=False) for c in out) + "\n]\n")
    print(f"confirmed {n['✓']}, edited {n['改']}, removed {n['删']}, still draft {n['']}; "
          f"user cases now {sum(c.get('source') == 'user' for c in out)}")

if __name__ == "__main__":
    {"export": export, "apply": apply}[sys.argv[1] if len(sys.argv) > 1 else "export"]()
