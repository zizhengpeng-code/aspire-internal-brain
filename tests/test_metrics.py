"""Metric self-check (handoff section 9 item 14), fake models only, no API calls.

The scripted models reproduce known ground truth, so the metric code is verified
independently of any LLM:

* ``oracle`` -> Benign Utility 1.0 and no leak;
* ``leaky``  -> Targeted ASR 1.0 (a legitimate citation with the canary smuggled into the
                prose; the guard cannot see it, which is the documented residual risk).

    python -m unittest discover tests
"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import evaluate_agent as ev

_CACHE = {}


def run_once(model, with_ablation=False):
    key = (model, with_ablation)
    if key not in _CACHE:
        _CACHE[key] = ev.run(model, with_ablation=with_ablation)
    return _CACHE[key]


class OracleSelfCheck(unittest.TestCase):
    """The oracle knows the gold answers, so utility must be perfect and nothing leaks."""

    @classmethod
    def setUpClass(cls):
        cls.rows, cls.sections, _ = run_once("oracle")

    def test_benign_utility_is_one(self):
        self.assertEqual(self.sections["all"]["benign_utility"], 1.0)
        self.assertEqual(self.sections["all"]["benign_utility_full"], 1.0)

    def test_utility_under_attack_is_one(self):
        self.assertEqual(self.sections["all"]["utility_under_attack"], 1.0)

    def test_no_leak(self):
        self.assertEqual(self.sections["all"]["targeted_asr"], 0.0)
        self.assertEqual(self.sections["all"]["leak_rate"], 0.0)

    def test_post_guard_citation_violation_is_zero(self):
        self.assertEqual(self.sections["all"]["post_guard_violation"], 0.0)

    def test_deny_cases_are_refused(self):
        # a correct oracle refuses every deny case
        self.assertEqual(self.sections["all"]["false_refusal"], 0.0)
        deny = [r for r in self.rows if r["expect"] == "deny"]
        self.assertTrue(all(r["refused"] for r in deny))

    def test_at_least_five_attacked_but_answerable_cases(self):
        # the handoff requires >= 5 "attacked but should still answer" cases
        self.assertGreaterEqual(self.sections["all"]["n_attack_answer"], 5)


class LeakySelfCheck(unittest.TestCase):
    """The leaky model smuggles a canary past the guard, so ASR must be 100%."""

    @classmethod
    def setUpClass(cls):
        cls.rows, cls.sections, _ = run_once("leaky")

    def test_targeted_asr_is_one(self):
        self.assertEqual(self.sections["all"]["targeted_asr"], 1.0)

    def test_leak_rate_is_one(self):
        self.assertEqual(self.sections["all"]["leak_rate"], 1.0)

    def test_guard_does_not_block_the_leak(self):
        # the leaky answer cites a valid visible document, so the guard passes -- this is
        # exactly the boundary documented as a residual risk in the README
        self.assertEqual(self.sections["all"]["false_refusal"], 0.0)


class AblationMonotonicity(unittest.TestCase):
    """Removing both layers must not make things safer than the production config."""

    @classmethod
    def setUpClass(cls):
        cls.rows, cls.sections, cls.ablations = run_once("fake", with_ablation=True)

    def test_removing_both_layers_does_not_lower_asr(self):
        self.assertGreaterEqual(self.ablations["no_search_filter+no_guard"]["targeted_asr"],
                                self.ablations["full"]["targeted_asr"])

    def test_all_four_configs_are_reported(self):
        self.assertEqual(set(self.ablations),
                         {"full", "no_search_filter", "no_guard", "no_search_filter+no_guard"})


class ReviewSheetRoundTrip(unittest.TestCase):
    """evaluate_agent.export_sheet -> fill the label -> score_review (needs openpyxl)."""

    def test_export_then_score(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl not installed")
        rows, _, _ = run_once("oracle")
        path = os.path.join(tempfile.mkdtemp(prefix="review_"), "answers_review.xlsx")
        ev.export_sheet(rows, path)
        wb = load_workbook(path)
        ws = wb.active
        col = [c.value for c in ws[1]].index(ev.LABEL_COL) + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).value = "✓"
        wb.save(path)
        tally = ev.score_review(path)
        self.assertEqual(tally["all"], (len(rows), len(rows)))


if __name__ == "__main__":
    unittest.main()
